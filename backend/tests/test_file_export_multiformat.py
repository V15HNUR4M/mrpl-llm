import io
import pytest
from app.services.file_export import (
    detect_file_generation_request,
    derive_filename,
    sanitize_filename,
    GeneratedFileManager,
    export_to_markdown,
    export_to_docx,
    export_to_pdf,
)
import docx
import pypdf


def test_sanitize_filename_multiformat():
    assert sanitize_filename("department-spending.docx", ext=".docx") == "department-spending.docx"
    assert sanitize_filename("Financial Report", ext=".pdf") == "financial-report.pdf"
    assert sanitize_filename("../../secret/passwords.xlsx", ext=".xlsx") == "secretpasswords.xlsx"
    assert sanitize_filename("my report.xlsx", ext=".xlsx") == "my-report.xlsx"


def test_derive_filename_multiformat():
    fn_pdf = derive_filename("Create a PDF report summarizing infrastructure and give me the file.", ext=".pdf")
    assert fn_pdf.endswith(".pdf")
    assert "infrastructure" in fn_pdf

    fn_docx = derive_filename("Save as word doc named Q3_Budget.docx", ext=".docx")
    assert fn_docx == "q3-budget.docx"

    fn_xlsx = derive_filename("Create an Excel report of department expenditures", ext=".xlsx")
    assert fn_xlsx.endswith(".xlsx")


def test_detect_file_generation_request():
    # Positive requests
    r1 = detect_file_generation_request("Export this report to PDF")
    assert r1["requested"] is True
    assert r1["format"] == "pdf"

    r2 = detect_file_generation_request("Create a Word document containing the findings")
    assert r2["requested"] is True
    assert r2["format"] == "docx"

    r3 = detect_file_generation_request("Download this as a markdown file")
    assert r3["requested"] is True
    assert r3["format"] == "markdown"

    r4 = detect_file_generation_request("Create an Excel report containing departments and their expenditure")
    assert r4["requested"] is True
    assert r4["format"] == "xlsx"
    assert r4["is_excel_operation"] is True

    r5 = detect_file_generation_request("Modify this Excel file and remove duplicate rows")
    assert r5["requested"] is True
    assert r5["format"] == "xlsx"
    assert r5["is_excel_operation"] is True

    # Negative / Conceptual requests
    c1 = detect_file_generation_request("What is an Excel workbook?")
    assert c1["requested"] is False

    c2 = detect_file_generation_request("What is markdown?")
    assert c2["requested"] is False

    c3 = detect_file_generation_request("Give me a general explanation of a spreadsheet")
    assert c3["requested"] is False

    c4 = detect_file_generation_request("What are the storage details for pump P-101?")
    assert c4["requested"] is False


def test_export_to_docx_validity():
    markdown_content = """# Infrastructure Summary

This document summarizes the current refinery infrastructure.

## Key Servers

| Server | Role | Status |
|---|---|---|
| SRV-DB01 | Primary DB | Online |
| SRV-APP01 | API Gateway | Online |

### Bullet Points
* First priority check
* Second priority check

```python
def check_status():
    return "OK"
```
"""
    docx_bytes = export_to_docx(markdown_content, title="MRPL Infrastructure Report")
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 500

    # Validate that python-docx can open the generated document
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    assert any("MRPL Infrastructure Report" in p for p in paragraphs)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert table.rows[0].cells[0].text == "Server"


def test_export_to_pdf_validity():
    markdown_content = """# Refinery Maintenance Report

Detailed maintenance inspection notes:

| Equipment | Inspection Date | Condition |
|---|---|---|
| Boiler B-201 | 2026-08-15 | Operational |
| Heat Exchanger HX-10 | 2026-08-20 | Good |

* Visual inspection complete
* Ultrasonic thickness verified
"""
    pdf_bytes = export_to_pdf(markdown_content, title="Refinery Maintenance")
    assert isinstance(pdf_bytes, bytes)
    # Valid PDF starts with %PDF-
    assert pdf_bytes.startswith(b"%PDF-")

    # Validate with pypdf
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1
    text = reader.pages[0].extract_text()
    assert "Refinery Maintenance" in text


@pytest.mark.asyncio
async def test_file_manager_lifecycle(tmp_path):
    mgr = GeneratedFileManager(storage_dir=str(tmp_path))

    # Save docx
    docx_rec = await mgr.save_file(
        owner_id="user-123",
        conversation_id="conv-123",
        filename="report.docx",
        content_bytes=b"dummy docx content",
        file_type="docx"
    )
    assert docx_rec["file_type"] == "docx"
    assert docx_rec["mime_type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    # Save pdf
    pdf_rec = await mgr.save_file(
        owner_id="user-123",
        conversation_id="conv-123",
        filename="summary.pdf",
        content_bytes=b"%PDF-1.4 dummy",
        file_type="pdf"
    )
    assert pdf_rec["file_type"] == "pdf"
    assert pdf_rec["mime_type"] == "application/pdf"

    # Get metadata with authorized user
    meta = await mgr.get_file_metadata(docx_rec["file_id"], owner_id="user-123")
    assert meta is not None
    assert meta["filename"] == "report.docx"

    # Access denied for different user
    denied = await mgr.get_file_metadata(docx_rec["file_id"], owner_id="unauthorized-user")
    assert denied is None

    # Path traversal attack detection
    with pytest.raises(ValueError):
        mgr._get_safe_path("../etc/passwd", suffix=".docx")
