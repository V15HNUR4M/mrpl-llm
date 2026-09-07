import pytest
import openpyxl
from pathlib import Path
import io

from app.core.agent.registry import AgentRegistry, AgentDefinition
from app.core.context_engine.engine import ContextEngine
from app.core.model_gateway.gateway import ModelGateway
from app.core.model_gateway.schemas import GenerationRequest, GenerationResponse, Usage
from app.core.model_gateway.provider import ModelProvider
from app.core.runtime.tool_executor import ToolRegistry, LocalToolExecutor, AuthorizedToolExecutor
from app.core.runtime.harness import AgentHarness
from app.core.runtime.schemas import RuntimeSession, ExecutionState
from app.core.excel.tools import (
    CreateWorkbookTool,
    SaveWorkbookTool,
    ReadWorkbookTool,
    InspectSheetTool,
    AggregateDataTool,
    FilterRowsTool,
    DetectDuplicatesTool,
    DetectMissingValuesTool,
    WriteCellsTool,
    AddFormulaTool,
    FormatRangeTool,
)
from app.services.file_export import GeneratedFileManager


class MockExcelLLMProvider(ModelProvider):
    """
    Simulates the LLM's response when asked to generate a spreadsheet with sample data.
    """
    def __init__(self):
        self.call_count = 0

    async def generate(self, request: GenerationRequest) -> GenerationResponse:
        self.call_count += 1
        if self.call_count == 1:
            # First turn: LLM generates the requested tabular dataset
            table_markdown = (
                "| Department | Employee Count | Budget in Lakhs |\n"
                "|---|---|---|\n"
                "| Refining | 240 | 185.5 |\n"
                "| Mechanical | 130 | 95.0 |\n"
                "| Electrical | 85 | 62.4 |\n"
                "| Safety | 45 | 38.0 |\n\n"
                "The new workbook with the added data is saved."
            )
            return GenerationResponse(text=table_markdown, model=request.model, usage=Usage())
        else:
            # Second turn: LLM receives deterministic tool execution output and returns confirmed answer
            return GenerationResponse(
                text="The Excel spreadsheet has been generated with 3 columns and 4 rows of data and saved.",
                model=request.model,
                usage=Usage()
            )

    async def stream(self, request: GenerationRequest):
        res = await self.generate(request)
        from app.core.model_gateway.schemas import StreamingEvent
        yield StreamingEvent(type="text_delta", text=res.text)
        yield StreamingEvent(type="completed", usage=res.usage)

    async def health(self): return True
    async def get_model_info(self, model_id): pass
    async def list_models(self): pass


@pytest.fixture
def excel_harness_setup(tmp_path, monkeypatch):
    from app.core.config import settings
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path / "uploads"))
    storage_dir = GeneratedFileManager().storage_dir

    tool_reg = ToolRegistry()
    tool_reg.register(CreateWorkbookTool())
    tool_reg.register(SaveWorkbookTool())
    tool_reg.register(ReadWorkbookTool())
    tool_reg.register(InspectSheetTool())
    tool_reg.register(AggregateDataTool())
    tool_reg.register(FilterRowsTool())
    tool_reg.register(DetectDuplicatesTool())
    tool_reg.register(DetectMissingValuesTool())
    tool_reg.register(WriteCellsTool())
    tool_reg.register(AddFormulaTool())
    tool_reg.register(FormatRangeTool())

    local_exec = LocalToolExecutor(tool_reg)
    auth_exec = AuthorizedToolExecutor(local_exec)

    agent_reg = AgentRegistry()
    excel_agent = AgentDefinition(
        agent_id="excel_agent",
        name="Excel Agent",
        description="Specialized agent for Excel operations",
        version="1.0",
        instructions=(
            "You are the MRPL Excel Specialist. Your workflow is: "
            "Inspect → execute deterministic tool operations → verify results → provide concise evidence-based answer. "
            "When a task requires creating or generating a workbook, you MUST invoke the create_workbook tool with clean headers and rows."
        ),
        tool_permissions={
            "allowed": [
                "create_workbook",
                "save_workbook",
                "read_workbook",
                "inspect_sheet",
                "aggregate_data",
                "filter_rows",
                "detect_duplicates",
                "detect_missing_values",
                "write_cells",
                "add_formula",
                "format_range",
            ]
        }
    )
    agent_reg.register(excel_agent)

    context_engine = ContextEngine()
    gateway = ModelGateway()
    mock_provider = MockExcelLLMProvider()
    gateway.register_provider("mock", mock_provider)
    from app.core.config import settings
    gateway.register_model_route(settings.DEFAULT_CHAT_MODEL, "mock")

    harness = AgentHarness(agent_reg, context_engine, gateway, auth_exec)
    return harness, storage_dir


@pytest.mark.asyncio
async def test_excel_agent_creates_downloadable_xlsx_end_to_end(excel_harness_setup):
    """
    Exact User Request:
    'Create an Excel spreadsheet with 3 columns: Department, Employee Count, and Budget in Lakhs. Add 4 rows of sample data.'

    Verifies:
    1. Excel Agent selects the spreadsheet-generation workflow.
    2. create_workbook() is invoked.
    3. An actual .xlsx file is produced.
    4. The file exists in the staging/storage location.
    5. The workbook contains 3 columns and 4 data rows.
    6. The generated workbook contains no Markdown table syntax.
    7. A valid file/download reference is returned.
    """
    harness, storage_dir = excel_harness_setup
    session = RuntimeSession(
        session_id="session-excel-001",
        conversation_id="conv-excel-001",
        agent_id="excel_agent",
        agent_version="1.0",
        user_id="analyst-user-42"
    )

    user_prompt = "Create an Excel spreadsheet with 3 columns: Department, Employee Count, and Budget in Lakhs. Add 4 rows of sample data."

    decision = await harness.execute(session, user_prompt)

    # 1. Verify Excel Agent completed execution
    assert session.state == ExecutionState.COMPLETED
    assert session.iteration_count >= 2

    # 2. Verify create_workbook tool was invoked
    tool_completed_events = [
        e for e in session.events
        if e.type == "tool_completed" and e.data.get("tool") == "create_workbook"
    ]
    assert len(tool_completed_events) == 1, "create_workbook tool was not invoked by Excel Agent"
    
    tool_result = tool_completed_events[0].data.get("result", {})
    assert tool_result.get("status") == "created_and_saved"
    assert tool_result.get("total_rows") == 4
    assert tool_result.get("total_columns") == 3

    # 3. Verify valid file/download reference is returned
    file_id = tool_result.get("file_id")
    filename = tool_result.get("filename")
    assert file_id is not None, "Missing file_id in tool result"
    assert filename.endswith(".xlsx"), f"Invalid filename {filename}"
    assert tool_result.get("file_type") == "xlsx"
    assert tool_result.get("download_url") == f"/api/v1/files/download/{file_id}"

    # 4. Verify file exists in storage location
    expected_file_path = storage_dir / f"{file_id}.xlsx"
    assert expected_file_path.exists(), f"Generated .xlsx file missing at {expected_file_path}"
    assert expected_file_path.stat().st_size > 0

    # 5. Open and inspect workbook with openpyxl
    wb = openpyxl.load_workbook(str(expected_file_path))
    ws = wb.active

    # Check 3 columns in header row
    headers = [ws.cell(row=1, column=c).value for c in range(1, 4)]
    assert headers == ["Department", "Employee Count", "Budget in Lakhs"]

    # Check 4 data rows
    data_rows = []
    for r in range(2, 6):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 4)]
        data_rows.append(row_vals)

    assert len(data_rows) == 4, f"Expected 4 data rows, found {len(data_rows)}"
    assert data_rows[0][0] == "Refining"
    assert data_rows[1][0] == "Mechanical"
    assert data_rows[2][0] == "Electrical"
    assert data_rows[3][0] == "Safety"

    # Check numeric types (not strings)
    assert isinstance(data_rows[0][1], (int, float)), f"Expected numeric employee count, got {type(data_rows[0][1])}"
    assert isinstance(data_rows[0][2], (int, float)), f"Expected numeric budget, got {type(data_rows[0][2])}"
    assert data_rows[0][1] == 240
    assert data_rows[0][2] == 185.5

    # 6. Verify NO Markdown table syntax inside any cell
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is not None:
                cell_str = str(cell)
                assert "|" not in cell_str, f"Cell contains markdown delimiter: {cell_str}"
                assert "---" not in cell_str, f"Cell contains markdown divider: {cell_str}"

    wb.close()


@pytest.mark.asyncio
async def test_excel_agent_stream_execute_workflow(excel_harness_setup):
    """
    Verifies that stream_execute also invokes create_workbook and delivers tool events.
    """
    harness, storage_dir = excel_harness_setup
    session = RuntimeSession(
        session_id="session-stream-002",
        conversation_id="conv-stream-002",
        agent_id="excel_agent",
        agent_version="1.0",
        user_id="analyst-user-42"
    )

    user_prompt = "Create an Excel spreadsheet with 3 columns: Department, Employee Count, and Budget in Lakhs. Add 4 rows of sample data."

    events = []
    async for ev in harness.stream_execute(session, user_prompt):
        events.append(ev)

    # Verify tool_completed event was emitted during stream
    tool_events = [e for e in events if e.get("type") == "tool_completed" and e.get("tool") == "create_workbook"]
    assert len(tool_events) == 1

    tool_res = tool_events[0].get("result", {})
    file_id = tool_res.get("file_id")
    assert file_id is not None

    xlsx_path = storage_dir / f"{file_id}.xlsx"
    assert xlsx_path.exists()


@pytest.mark.asyncio
async def test_excel_agent_api_endpoint_workflow(async_client):
    """
    Full API-level test:
    POST /api/v1/agents/excel_agent/run -> create_workbook -> GET /api/v1/files/download/{file_id}
    """
    from app.main import app
    from app.db.database import AsyncSessionLocal
    from app.db.uow import UnitOfWork
    from app.services.auth import AuthService
    import uuid

    uow = UnitOfWork(session_factory=AsyncSessionLocal)
    user_id = str(uuid.uuid4())
    username = f"excel_user_{user_id[:8]}"
    password = "SecretPassword123!"

    from app.core.security import create_access_token

    async with uow:
        auth_service = AuthService(uow)
        user = await auth_service.create_user({
            "username": username,
            "password": password,
            "email": f"{username}@example.com",
            "role": "USER",
            "is_active": True
        })
        token = create_access_token(data={"sub": user.username, "role": user.role})
        
        conv = await uow.conversations.create({
            "id": str(uuid.uuid4()),
            "user_id": user.id,
            "title": "Excel Generation Test"
        })
        await uow.commit()

    headers = {"Authorization": f"Bearer {token}"}

    # Register MockExcelLLMProvider on app gateway
    mock_provider = MockExcelLLMProvider()
    app.state.model_gateway.register_provider("mock_excel_api", mock_provider)
    from app.core.config import settings
    app.state.model_gateway.register_model_route(settings.DEFAULT_CHAT_MODEL, "mock_excel_api")

    user_prompt = "Create an Excel spreadsheet with 3 columns: Department, Employee Count, and Budget in Lakhs. Add 4 rows of sample data."

    run_res = await async_client.post(
        "/api/v1/agents/excel_agent/run",
        json={
            "conversation_id": conv.id,
            "message": user_prompt
        },
        headers=headers
    )

    assert run_res.status_code == 200, f"Run failed: {run_res.text}"
    data = run_res.json()
    assert "events" in data
    
    # Check that tool_completed event for create_workbook is present
    tool_events = [e for e in data["events"] if e.get("type") == "tool_completed" and e.get("data", {}).get("tool") == "create_workbook"]
    assert len(tool_events) == 1, "create_workbook was not invoked"

    gen_file = data.get("generated_file")
    assert gen_file is not None, "Missing generated_file in API response"
    file_id = gen_file.get("file_id")
    assert file_id is not None
    assert gen_file.get("file_type") == "xlsx"

    # Now verify the download endpoint works
    download_res = await async_client.get(
        f"/api/v1/files/download/{file_id}",
        headers=headers
    )
    assert download_res.status_code == 200
    assert "spreadsheetml" in download_res.headers.get("content-type", "")

    # Open downloaded content with openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(download_res.content))
    ws = wb.active

    # Check 3 columns
    col_headers = [ws.cell(row=1, column=c).value for c in range(1, 4)]
    assert col_headers == ["Department", "Employee Count", "Budget in Lakhs"]

    # Check 4 rows
    rows_data = []
    for r in range(2, 6):
        rows_data.append([ws.cell(row=r, column=c).value for c in range(1, 4)])

    assert len(rows_data) == 4
    assert rows_data[0][0] == "Refining"
    assert rows_data[1][0] == "Mechanical"
    assert rows_data[2][0] == "Electrical"
    assert rows_data[3][0] == "Safety"

    # Verify no markdown delimiters in cells
    for row in ws.iter_rows(values_only=True):
        for cell in row:
            if cell is not None:
                assert "|" not in str(cell)
                assert "---" not in str(cell)

    wb.close()

