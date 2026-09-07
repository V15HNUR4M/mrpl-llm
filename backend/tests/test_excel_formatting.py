import pytest
import openpyxl
from pathlib import Path
from app.core.excel.table_parser import (
    clean_cell_str,
    coerce_value,
    parse_markdown_table,
    extract_tabular_data,
)
from app.core.excel.writers import create_workbook
from app.services.file_export import GeneratedFileManager


SAMPLE_MARKDOWN_TABLE = """
Here is the requested employee and budget breakdown:

| Department | Employee Count | Budget (Lakhs) |
| --- | --- | --- |
| Sales | 10 | 5 |
| Marketing | 12 | 8 |
| IT | 8 | 3 |
| HR | 15 | 10 |

Please let me know if you need any adjustments.
"""


def test_markdown_table_parsing():
    headers, rows = parse_markdown_table(SAMPLE_MARKDOWN_TABLE)

    assert headers == ["Department", "Employee Count", "Budget (Lakhs)"]
    assert len(rows) == 4

    assert rows[0] == ["Sales", 10, 5]
    assert rows[1] == ["Marketing", 12, 8]
    assert rows[2] == ["IT", 8, 3]
    assert rows[3] == ["HR", 15, 10]

    # Verify numeric types
    for r in rows:
        assert isinstance(r[1], (int, float))
        assert isinstance(r[2], (int, float))
        assert not any("|" in str(cell) for cell in r)
        assert not any("---" in str(cell) for cell in r)


def test_extract_tabular_data_fallback_and_cleaning():
    # Test table with bold markdown syntax inside cells
    bold_md = """
    | **Department** | **Employee Count** | **Budget (Lakhs)** |
    | :--- | :---: | ---: |
    | `Sales` | 10 | 5.50 |
    | *Marketing* | 12 | 8.25 |
    """
    headers, rows = extract_tabular_data(bold_md)
    assert headers == ["Department", "Employee Count", "Budget (Lakhs)"]
    assert rows[0] == ["Sales", 10, 5.5]
    assert rows[1] == ["Marketing", 12, 8.25]
    assert isinstance(rows[0][2], float)


def test_create_workbook_formatting_and_openpyxl_inspection():
    headers, rows = parse_markdown_table(SAMPLE_MARKDOWN_TABLE)

    res = create_workbook(sheet_name="Data", headers=headers, rows=rows)
    staging_id = res["staging_file_id"]
    assert res["status"] == "created"
    assert res["total_rows"] == 4
    assert res["total_columns"] == 3

    # Programmatically inspect the generated .xlsx using openpyxl
    file_mgr = GeneratedFileManager()
    file_path = file_mgr.storage_dir / f"{staging_id}.xlsx"
    assert file_path.exists()

    wb = openpyxl.load_workbook(str(file_path))
    assert wb.sheetnames == ["Data"]
    ws = wb["Data"]

    # 1. Correct number of rows and columns
    assert ws.max_row == 5  # 1 header + 4 data rows
    assert ws.max_column == 3

    # 2. Headers are separate cells, properly styled and free of markdown syntax
    header_cells = [ws.cell(row=1, column=c).value for c in range(1, 4)]
    assert header_cells == ["Department", "Employee Count", "Budget (Lakhs)"]

    for c in range(1, 4):
        cell = ws.cell(row=1, column=c)
        assert cell.font.bold is True
        assert cell.font.color.rgb == "00FFFFFF" or cell.font.color.rgb == "FFFFFF"
        assert cell.fill.start_color.rgb == "000F172A" or cell.fill.start_color.rgb == "0F172A"
        assert cell.border.left.style == "thin"
        assert cell.border.bottom.style == "thin"

    # Header row height
    assert ws.row_dimensions[1].height == 26

    # 3. Freeze row applied below header
    assert ws.freeze_panes == "A2"

    # 4. Auto-filter enabled across table range
    assert ws.auto_filter.ref == "A1:C5"

    # 5. Data rows validation
    expected_data = [
        ["Sales", 10, 5],
        ["Marketing", 12, 8],
        ["IT", 8, 3],
        ["HR", 15, 10]
    ]

    for r_idx, exp_row in enumerate(expected_data, start=2):
        assert ws.row_dimensions[r_idx].height == 20
        # Text column (Department)
        dept_cell = ws.cell(row=r_idx, column=1)
        assert dept_cell.value == exp_row[0]
        assert dept_cell.alignment.horizontal == "left"
        assert dept_cell.border.left.style == "thin"

        # Numeric column (Employee Count)
        emp_cell = ws.cell(row=r_idx, column=2)
        assert emp_cell.value == exp_row[1]
        assert isinstance(emp_cell.value, int)
        assert emp_cell.alignment.horizontal == "right"
        assert emp_cell.border.left.style == "thin"

        # Numeric column (Budget)
        budget_cell = ws.cell(row=r_idx, column=3)
        assert budget_cell.value == exp_row[2]
        assert isinstance(budget_cell.value, int)
        assert budget_cell.alignment.horizontal == "right"
        assert budget_cell.border.left.style == "thin"

    # 6. Column widths
    assert ws.column_dimensions["A"].width >= 14
    assert ws.column_dimensions["B"].width >= 14
    assert ws.column_dimensions["C"].width >= 14

    wb.close()
