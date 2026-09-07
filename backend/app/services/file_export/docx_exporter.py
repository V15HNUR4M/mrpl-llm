import io
import re
from typing import Optional, List
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn


def _set_cell_background(cell, fill_hex: str):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def _add_formatted_text(paragraph, text: str):
    """
    Parses basic markdown bold and italic inline tokens into docx runs.
    """
    # Regex splits by **bold** or *italic*
    tokens = re.split(r'(\*\*[^*]+?\*\*|\*[^*]+?\*)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)


def export_to_docx(content: str, title: Optional[str] = None) -> bytes:
    """
    Converts markdown content into a professionally styled Microsoft Word document (.docx).
    Handles headings, tables, bullet lists, numbered lists, and code blocks.
    Returns bytes of the docx file.
    """
    doc = Document()

    # Document-wide margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Optional Title
    if title:
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(title)
        title_run.font.size = Pt(22)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(15, 23, 42) # slate-900
        title_p.paragraph_format.space_after = Pt(12)

    lines = content.splitlines()
    i = 0
    in_code_block = False
    code_lines: List[str] = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 1. Code Block start/end
        if stripped.startswith("```"):
            if in_code_block:
                # Flush code block
                code_text = "\n".join(code_lines)
                code_p = doc.add_paragraph()
                code_p.paragraph_format.left_indent = Inches(0.25)
                code_p.paragraph_format.space_before = Pt(4)
                code_p.paragraph_format.space_after = Pt(6)
                
                run = code_p.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
                
                # Background shading for paragraph
                pPr = code_p._p.get_or_add_pPr()
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
                pPr.append(shd)

                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Blank lines
        if not stripped:
            i += 1
            continue

        # 2. Markdown Table
        if stripped.startswith("|") and stripped.endswith("|"):
            table_rows = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                row_str = lines[i].strip()
                # Skip separator row e.g. |---|---|
                if re.match(r"^\|(\s*:?-+:?\s*\|)+$", row_str):
                    i += 1
                    continue
                cells = [c.strip() for c in row_str.strip("|").split("|")]
                table_rows.append(cells)
                i += 1

            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
                tbl.autofit = True

                for r_idx, row_data in enumerate(table_rows):
                    row = tbl.rows[r_idx]
                    is_header = (r_idx == 0)
                    for c_idx in range(num_cols):
                        cell = row.cells[c_idx]
                        cell_val = row_data[c_idx] if c_idx < len(row_data) else ""
                        cell.text = ""
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)

                        if is_header:
                            _set_cell_background(cell, "0F172A")
                            run = p.add_run(cell_val)
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            run.font.size = Pt(10)
                        else:
                            if r_idx % 2 == 1:
                                _set_cell_background(cell, "F8FAFC")
                            run = p.add_run(cell_val)
                            run.font.size = Pt(9.5)
                            run.font.color.rgb = RGBColor(51, 65, 85)
                
                # Add spacing after table
                space_p = doc.add_paragraph()
                space_p.paragraph_format.space_after = Pt(6)
            continue

        # 3. Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            run = h.add_run(stripped[2:].strip())
            run.font.color.rgb = RGBColor(15, 23, 42)
            run.font.size = Pt(18)
            i += 1
            continue

        if stripped.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            run = h.add_run(stripped[3:].strip())
            run.font.color.rgb = RGBColor(30, 41, 59)
            run.font.size = Pt(14)
            i += 1
            continue

        if stripped.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            run = h.add_run(stripped[4:].strip())
            run.font.color.rgb = RGBColor(51, 65, 85)
            run.font.size = Pt(12)
            i += 1
            continue

        # 4. Bullet lists (* or -)
        if re.match(r"^[\*\-]\s+", stripped):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            text_part = re.sub(r"^[\*\-]\s+", "", stripped)
            _add_formatted_text(p, text_part)
            i += 1
            continue

        # 5. Numbered lists (1. )
        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            text_part = re.sub(r"^\d+\.\s+", "", stripped)
            _add_formatted_text(p, text_part)
            i += 1
            continue

        # 6. Standard Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        _add_formatted_text(p, stripped)
        i += 1

    # Save to buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
