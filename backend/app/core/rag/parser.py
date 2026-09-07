from abc import ABC, abstractmethod
from typing import Dict, Any, Type, List
import io
import pypdf

from app.core.rag.schemas import ParsedDocument
from app.core.rag.errors import UnsupportedFormatError, ParsingError

class DocumentParser(ABC):
    @classmethod
    @abstractmethod
    def supports(cls, mime_type: str, file_type: str) -> bool:
        pass

    @abstractmethod
    def parse(self, file_bytes: bytes, metadata: Dict[str, Any] = None) -> ParsedDocument:
        pass

class TXTParser(DocumentParser):
    @classmethod
    def supports(cls, mime_type: str, file_type: str) -> bool:
        return mime_type == "text/plain" or file_type.lower() == "txt"

    def parse(self, file_bytes: bytes, metadata: Dict[str, Any] = None) -> ParsedDocument:
        try:
            content = file_bytes.decode('utf-8')
            return ParsedDocument(content=content, metadata=metadata or {})
        except UnicodeDecodeError:
            try:
                # Fallback
                content = file_bytes.decode('latin-1')
                return ParsedDocument(content=content, metadata=metadata or {})
            except Exception as e:
                raise ParsingError(f"Failed to decode TXT file: {str(e)}")

class MarkdownParser(DocumentParser):
    @classmethod
    def supports(cls, mime_type: str, file_type: str) -> bool:
        return mime_type in ["text/markdown", "text/x-markdown"] or file_type.lower() == "md"

    def parse(self, file_bytes: bytes, metadata: Dict[str, Any] = None) -> ParsedDocument:
        try:
            content = file_bytes.decode('utf-8')
            return ParsedDocument(content=content, metadata=metadata or {})
        except Exception as e:
            raise ParsingError(f"Failed to decode Markdown file: {str(e)}")

class PDFParser(DocumentParser):
    @classmethod
    def supports(cls, mime_type: str, file_type: str) -> bool:
        return mime_type == "application/pdf" or file_type.lower() == "pdf"

    def parse(self, file_bytes: bytes, metadata: Dict[str, Any] = None) -> ParsedDocument:
        try:
            pdf_file = io.BytesIO(file_bytes)
            reader = pypdf.PdfReader(pdf_file)
            
            pages_text = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    pages_text.append(f"--- Page {i+1} ---\n{text}")
            
            content = "\n\n".join(pages_text)
            return ParsedDocument(content=content, metadata=metadata or {})
        except Exception as e:
            raise ParsingError(f"Failed to parse PDF file: {str(e)}")

class XLSXParser(DocumentParser):
    @classmethod
    def supports(cls, mime_type: str, file_type: str) -> bool:
        xlsx_mimes = [
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel"
        ]
        return mime_type in xlsx_mimes or file_type.lower() in ["xlsx", "xls"]

    def parse(self, file_bytes: bytes, metadata: Dict[str, Any] = None) -> ParsedDocument:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            sheet_texts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows_text = [f"### Sheet: {sheet_name}"]
                for row_idx, row in enumerate(ws.iter_rows(values_only=True, max_row=200)):
                    row_vals = [str(c) if c is not None else "" for c in row]
                    if any(row_vals):
                        rows_text.append(" | ".join(row_vals))
                sheet_texts.append("\n".join(rows_text))
            wb.close()
            content = "\n\n".join(sheet_texts)
            return ParsedDocument(content=content, metadata=metadata or {})
        except Exception as e:
            raise ParsingError(f"Failed to parse Excel file: {str(e)}")

class ParserRegistry:
    def __init__(self):
        self._parsers: List[Type[DocumentParser]] = [
            TXTParser,
            MarkdownParser,
            PDFParser,
            XLSXParser
        ]

    def get_parser(self, mime_type: str, file_type: str) -> DocumentParser:
        for parser_cls in self._parsers:
            if parser_cls.supports(mime_type, file_type):
                return parser_cls()
        raise UnsupportedFormatError(f"No parser supports mime_type={mime_type}, file_type={file_type}")

